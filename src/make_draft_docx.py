"""Markdown draft -> .docx, for uploading to Google Drive as a Doc.

Handles the subset used in the drafts: headings, paragraphs, bullet and numbered
lists, tables, block quotes, fenced code, images, and inline bold/italic/code/links.

usage: python3 src/make_draft_docx.py 8
"""
import re, sys
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).resolve().parent.parent
V = sys.argv[1] if len(sys.argv) > 1 else "8"
SRC = ROOT / f"DRAFT_writeup_v{V}.md"
OUT = ROOT / "writeup_pack" / f"DRAFT_v{V}.docx"

INLINE = re.compile(r"(\*\*.+?\*\*|\*[^*]+?\*|`[^`]+?`|\[[^\]]+\]\([^)]+\))", re.S)

def add_runs(par, text):
    for piece in INLINE.split(text):
        if not piece:
            continue
        if piece.startswith("**") and piece.endswith("**"):
            par.add_run(piece[2:-2]).bold = True
        elif piece.startswith("*") and piece.endswith("*") and len(piece) > 2:
            par.add_run(piece[1:-1]).italic = True
        elif piece.startswith("`") and piece.endswith("`"):
            r = par.add_run(piece[1:-1]); r.font.name = "Consolas"; r.font.size = Pt(9.5)
        elif piece.startswith("[") and re.match(r"\[([^\]]+)\]\(([^)]+)\)", piece):
            label, url = re.match(r"\[([^\]]+)\]\(([^)]+)\)", piece).groups()
            r = par.add_run(label); r.font.color.rgb = RGBColor(0x15, 0x5C, 0xB0)
            if url.startswith("http"):
                r2 = par.add_run(f" ({url})"); r2.font.size = Pt(8.5)
                r2.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
        else:
            par.add_run(piece.replace("\\*", "*").replace("\\_", "_"))

def main():
    lines = SRC.read_text().split("\n")
    doc = Document()
    st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(11)
    i = 0
    while i < len(lines):
        ln = lines[i]
        if ln.startswith("```"):                                    # code block
            i += 1; buf = []
            while i < len(lines) and not lines[i].startswith("```"):
                buf.append(lines[i]); i += 1
            p = doc.add_paragraph(); p.paragraph_format.left_indent = Inches(0.3)
            r = p.add_run("\n".join(buf)); r.font.name = "Consolas"; r.font.size = Pt(9)
            i += 1; continue
        if ln.startswith("|") and i + 1 < len(lines) and set(lines[i+1].replace("|", "").strip()) <= set("-: "):
            rows = []                                               # table
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not set("".join(cells)) <= set("-: "):
                    rows.append(cells)
                i += 1
            width = max(len(r) for r in rows)
            t = doc.add_table(rows=0, cols=width); t.style = "Table Grid"
            for j, row in enumerate(rows):
                cs = t.add_row().cells
                for k in range(width):
                    cell = cs[k]; cell.text = ""
                    par = cell.paragraphs[0]; add_runs(par, row[k] if k < len(row) else "")
                    for r in par.runs:
                        r.font.size = Pt(9.5)
                        if j == 0: r.bold = True
            doc.add_paragraph(); continue
        m = re.match(r"!\[([^\]]*)\]\(([^)]+)\)", ln.strip())        # image
        if m:
            path = ROOT / m.group(2)
            if path.exists():
                doc.add_picture(str(path), width=Inches(6.0))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1; continue
        if ln.startswith("#"):                                       # heading
            lvl = len(ln) - len(ln.lstrip("#"))
            doc.add_heading(ln.lstrip("#").strip(), level=min(lvl, 4)); i += 1; continue
        if ln.strip() == "---":
            doc.add_paragraph(); i += 1; continue
        if ln.startswith(">"):                                       # quote
            buf = []
            while i < len(lines) and lines[i].startswith(">"):
                buf.append(lines[i].lstrip("> ").rstrip()); i += 1
            p = doc.add_paragraph(); p.paragraph_format.left_indent = Inches(0.35)
            p.paragraph_format.space_before = Pt(6); p.paragraph_format.space_after = Pt(6)
            add_runs(p, " ".join(x for x in buf if x))
            for r in p.runs: r.italic = True
            continue
        mb = re.match(r"^(\s*)([-*]|\d+\.)\s+(.*)", ln)              # list item
        if mb:
            indent, marker, text = mb.groups()
            buf = [text]; i += 1
            while i < len(lines) and lines[i].startswith("  ") and lines[i].strip() and not re.match(r"^\s*([-*]|\d+\.)\s", lines[i]):
                buf.append(lines[i].strip()); i += 1
            style = "List Number" if marker[0].isdigit() else "List Bullet"
            p = doc.add_paragraph(style=style)
            if indent: p.paragraph_format.left_indent = Inches(0.25 + 0.25 * (len(indent) // 2))
            add_runs(p, " ".join(buf)); continue
        if not ln.strip():
            i += 1; continue
        buf = [ln]; i += 1                                           # paragraph
        while i < len(lines) and lines[i].strip() and not re.match(r"^(#|\||>|```|!\[|\s*([-*]|\d+\.)\s)", lines[i]):
            buf.append(lines[i]); i += 1
        add_runs(doc.add_paragraph(), " ".join(x.strip() for x in buf))
    doc.save(OUT)
    print("wrote", OUT)

main()
